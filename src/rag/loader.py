import io
from datetime import date
from pathlib import Path

import frontmatter
from langchain_core.documents import Document

from src.rag.config import KB_DIR

SUPPORTED_UPLOAD_EXTENSIONS = {".pdf", ".docx", ".md", ".txt"}


def load_kb_documents(kb_dir: Path = KB_DIR) -> list[Document]:
    """Load every KB markdown file into a LangChain Document, frontmatter as metadata."""
    documents = []
    for path in sorted(kb_dir.glob("*.md")):
        post = frontmatter.load(path)
        metadata = {
            "source_file": path.name,
            "article_id": post.get("article_id", path.stem),
            "title": post.get("title", path.stem),
            "product_area": post.get("product_area", "general"),
            "last_updated": str(post.get("last_updated", "")),
        }
        documents.append(Document(page_content=post.content, metadata=metadata))
    return documents


def _extract_pdf_text(data: bytes) -> str:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n\n".join(page.extract_text() or "" for page in pdf.pages)


def _extract_docx_text(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def load_uploaded_document(filename: str, data: bytes) -> Document:
    """Extract text from one uploaded file (pdf/docx/md/txt) into a Document.

    Raises ValueError for an unsupported extension or a file with no extractable text
    (e.g. a scanned/image-only PDF) so the caller can report it per-file instead of
    silently indexing an empty chunk.
    """
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf_text(data)
    elif suffix == ".docx":
        text = _extract_docx_text(data)
    elif suffix in (".md", ".txt"):
        text = data.decode("utf-8", errors="replace")
        if suffix == ".md":
            text = frontmatter.loads(text).content
    else:
        raise ValueError(
            f"Unsupported file type '{suffix}' (supported: {sorted(SUPPORTED_UPLOAD_EXTENSIONS)})"
        )

    if not text.strip():
        raise ValueError(f"No extractable text found in '{filename}'")

    metadata = {
        "source_file": filename,
        "article_id": Path(filename).stem,
        "title": Path(filename).stem,
        "product_area": "uploaded",
        "last_updated": str(date.today()),
    }
    return Document(page_content=text, metadata=metadata)


def load_uploaded_documents(files: list[tuple[str, bytes]]) -> tuple[list[Document], list[tuple[str, str]]]:
    """Load a batch of uploaded (filename, bytes) files.

    Returns (documents, failures) — failures is a list of (filename, error message)
    for files that couldn't be parsed, so the caller can report them individually
    instead of one upload failure aborting the whole batch.
    """
    documents, failures = [], []
    for filename, data in files:
        try:
            documents.append(load_uploaded_document(filename, data))
        except Exception as e:
            failures.append((filename, str(e)))
    return documents, failures
